"""Bounded, deterministic, format-aware digests of the resolved input files.

Purely descriptive: for every input this module reports what is in it (structure, sampled rows,
value/date ranges, categorical levels, the parameter values of small config files, the text of small
code/markdown files, the member list of workbooks/archives). It never computes anything that could
function as a reference answer and never opens anything outside the resolved input path it is given.

Two layers:
  * `summarize()` returns one JSON-serialisable dict per file. The structural keys (`format`,
    `n_rows`, `columns`, `dtypes`, `null_counts`, `shape`, `top_level_keys`, ...) are the original
    ones and are what `agent.self_check` reads; `n_rows` is still only set for csv/parquet, so the
    self-check's row-parity gate sees exactly the same tabular inputs as before. New detail lives in
    `sections`: size-ordered *variants* (largest first) of each piece of detail, so the prompt
    builder (`agent.digest`) can fit the global size budget by picking the largest variant that fits.
  * `cross_file_notes()` states how several tabular inputs relate (shared columns, overlap).

File contents are DATA. Everything handed to the prompt goes through `clean_text`: control
characters are dropped and UUID-shaped strings (the benchmark's canary GUIDs appear in several
input files) are replaced, so that a generated script cannot copy one out of the prompt.
Nothing here ever raises on a malformed file: a failure is recorded as `summary_error` and the
file still gets whatever generic text excerpt can be read.
"""

from __future__ import annotations

import json
import re
import warnings
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from agent.input_resolver import ResolvedInput

# --- Per-file size limits (characters of prompt text) ----------------------------------------------
# A file that the instruction itself names (by file name) is more likely to be task-critical, so its
# text/JSON content may use the larger cap. Everything is further bounded by the global budget in
# agent.digest.
TEXT_CAP = 6000
REFERENCED_TEXT_CAP = 20000
JSON_CAP = 8000
REFERENCED_JSON_CAP = 12000
_SHRUNK_CAPS = (2500, 800)  # the smaller variants offered to the global-budget packer

_MAX_FULL_READ_BYTES = 2_000_000  # a text file larger than this is read as head + tail only
_MAX_TABLE_BYTES = 200_000_000  # beyond this, tabular detail is computed from a sample
_SAMPLE_ROWS_HEAD = 5
_SAMPLE_ROWS_TAIL = 2
_MAX_SAMPLE_COLUMNS = 30
_MAX_NOTE_COLUMNS = 40
_CELL_CHARS = 40
_MAX_LEVELS = 12

# .pqt is a Parquet alias seen in this corpus (e.g. t1-double-sort, t1-residual-momentum,
# t1-stable-residual -- 3 units, actually more common here than the ".parquet" spelling itself,
# which appears in only 1). Same file format, same structured summary; dispatch on either
# spelling to the same parquet reader.
_PARQUET_EXTENSIONS = {".parquet", ".pqt"}
_HTML_EXTENSIONS = {".html", ".htm"}
_XML_EXTENSIONS = {".xml", ".xsd"}

_CONTROL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
# Hex-aware boundaries rather than \b: an identifier such as `loc_Revenues_<uuid>` has `_` (a word
# character) right before the UUID, which \b would not treat as a boundary.
_UUID_RE = re.compile(r"(?<![0-9a-fA-F])[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}(?![0-9a-fA-F])")
UUID_PLACEHOLDER = "<uuid-redacted>"
_DATE_LIKE_RE = re.compile(
    r"^\d{4}[-/]\d{2}[-/]\d{2}(?:[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?(?:Z|[+-]\d{2}:?\d{2})?)?$"
    r"|^\d{1,2}/\d{1,2}/\d{2,4}$"
)


def clean_text(text: str) -> str:
    """Make file-derived text safe to place in a prompt: no control characters, no UUID-shaped
    strings (see module docstring)."""
    return _UUID_RE.sub(UUID_PLACEHOLDER, _CONTROL_CHARS_RE.sub("", text))


def _section(kind: str, variants: list[dict[str, str]]) -> dict[str, Any]:
    return {"kind": kind, "variants": variants}


def _dedupe(variants: list[dict[str, str]]) -> list[dict[str, str]]:
    seen: set[str] = set()
    out = []
    for variant in variants:
        if variant["text"] not in seen:
            seen.add(variant["text"])
            out.append(variant)
    return out


def _is_referenced(resolved_input: ResolvedInput, instruction_text: str) -> bool:
    """True when the instruction names this file (by file name) -- a generic, task-agnostic signal
    that the file is part of the task statement rather than incidental data."""
    if not instruction_text:
        return False
    haystack = instruction_text.lower()
    return resolved_input.resolved_path.name.lower() in haystack


def summarize(resolved_input: ResolvedInput, instruction_text: str = "") -> dict[str, Any]:
    """Dispatch on file extension; never raises."""
    path = resolved_input.resolved_path
    suffix = path.suffix.lower()
    referenced = _is_referenced(resolved_input, instruction_text)

    base: dict[str, Any] = {
        "manifest_path": resolved_input.manifest_path,
        "filename": path.name,
        "bytes": path.stat().st_size,
        "referenced_in_instruction": referenced,
        "sections": [],
    }

    try:
        if suffix == ".csv":
            base.update(_summarize_csv(path))
        elif suffix in _PARQUET_EXTENSIONS:
            base.update(_summarize_parquet(path))
        elif suffix == ".tsv":
            base.update(_summarize_tsv(path))
        elif suffix == ".json":
            base.update(_summarize_json(path, referenced))
        elif suffix == ".jsonl":
            base.update(_summarize_jsonl(path))
        elif suffix == ".xlsx":
            base.update(_summarize_xlsx(path))
        elif suffix == ".zip":
            base.update(_summarize_zip(path))
        elif suffix in _HTML_EXTENSIONS:
            base.update(_summarize_html(path, referenced))
        elif suffix in _XML_EXTENSIONS:
            base.update(_summarize_xml(path, referenced))
        else:
            base.update(_summarize_text(path, referenced))
    except Exception as exc:  # noqa: BLE001 - a summary failure must not crash the pipeline
        base["summary_error"] = f"{type(exc).__name__}: {exc}"
        if not base["sections"]:
            try:
                base["sections"] = _text_sections(path, referenced) or []
            except Exception:  # noqa: BLE001
                pass

    return base


# ---------------------------------------------------------------------------------------------
# Tabular
# ---------------------------------------------------------------------------------------------


def _num(value: Any) -> str:
    x = float(value)
    if x != x:
        return "nan"
    if x.is_integer() and abs(x) < 1e15:
        return str(int(x))
    return f"{x:.6g}"


def _cell(value: Any) -> str:
    text = clean_text(str(value))
    return text if len(text) <= _CELL_CHARS else text[: _CELL_CHARS - 3] + "..."


def _date_note(parsed: Any, in_file_order_sorted: bool) -> str:
    """`parsed` is a non-null datetime Series."""
    lo, hi = parsed.min(), parsed.max()
    has_time = bool((parsed.dt.normalize() != parsed).any())
    fmt = "%Y-%m-%d %H:%M:%S" if has_time else "%Y-%m-%d"
    uniq = parsed.drop_duplicates().sort_values()
    step = ""
    if len(uniq) > 1:
        median_step = uniq.diff().dropna().median()
        step = f"; median step {median_step.total_seconds() / 86400:.4g} days"
    order = "ascending" if in_file_order_sorted else "not sorted ascending"
    return f"dates {lo.strftime(fmt)} to {hi.strftime(fmt)} ({len(uniq):,} distinct{step}; {order} in file order)"


def _series_note(series: Any) -> str | None:
    import pandas as pd
    from pandas.api import types as pt

    non_null = series.dropna()
    if len(non_null) == 0:
        return "all values missing"

    if pt.is_bool_dtype(series):
        return "bool; values " + ", ".join(sorted({str(v) for v in non_null.unique()}))

    if pt.is_datetime64_any_dtype(series):
        return _date_note(non_null, bool(non_null.is_monotonic_increasing))

    if pt.is_numeric_dtype(series):
        note = f"min={_num(non_null.min())}, max={_num(non_null.max())}, mean={_num(non_null.mean())}"
        distinct = non_null.nunique()
        if distinct <= 8:
            note += "; values " + ", ".join(sorted(_num(v) for v in non_null.unique()))
        return note

    texts = non_null.astype(str)
    sample = [t.strip() for t in texts.iloc[:50]]
    if sample and sum(bool(_DATE_LIKE_RE.match(t)) for t in sample) / len(sample) >= 0.9:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            parsed = pd.to_datetime(texts, errors="coerce", utc=True)
        if parsed.notna().mean() >= 0.9:
            parsed = parsed.dropna()
            return "date strings: " + _date_note(parsed, bool(parsed.is_monotonic_increasing))

    distinct = texts.nunique()
    if distinct <= _MAX_LEVELS:
        levels = sorted(texts.unique().tolist())
        return "levels: " + ", ".join(repr(_cell(v)) for v in levels)
    return f"{distinct:,} distinct strings"


def _tabular_sections(df: Any, label_suffix: str = "") -> list[dict[str, Any]]:
    """column_notes + sample sections for a DataFrame."""
    columns = list(df.columns)
    notes: list[str] = []
    for column in columns[:_MAX_NOTE_COLUMNS]:
        try:
            note = _series_note(df[column])
        except Exception:  # noqa: BLE001 - one odd column (duplicate name, nested values) must not matter
            note = None
        if note:
            notes.append(f"{_cell(column)}: {note}")
    if len(columns) > _MAX_NOTE_COLUMNS:
        notes.append(f"(+{len(columns) - _MAX_NOTE_COLUMNS} more columns not described)")

    sections: list[dict[str, Any]] = []
    if notes:
        short = notes[:12]
        if len(notes) > 12:
            short.append(f"(+{len(notes) - 12} more columns not described)")
        sections.append(
            _section(
                "column_notes",
                _dedupe(
                    [
                        {"label": f"per-column value summary{label_suffix}", "text": clean_text("\n".join(notes))},
                        {"label": f"per-column value summary (first 12 columns){label_suffix}", "text": clean_text("\n".join(short))},
                    ]
                ),
            )
        )

    shown = columns[:_MAX_SAMPLE_COLUMNS]
    omitted = len(columns) - len(shown)

    def render(frame: Any) -> str:
        frame = frame.iloc[:, :_MAX_SAMPLE_COLUMNS].copy()
        for column in frame.columns:
            if frame[column].dtype == object or str(frame[column].dtype) in {"str", "string"}:
                frame[column] = frame[column].map(_cell)
        text = frame.to_csv(index=False, float_format="%.8g")
        return clean_text(text.rstrip("\n"))

    def sample_text(n_head: int, n_tail: int) -> str:
        parts = [f"first {min(n_head, len(df))} rows:\n{render(df.iloc[:n_head])}"]
        if n_tail and len(df) > n_head + n_tail:
            parts.append(f"last {n_tail} rows:\n{render(df.iloc[-n_tail:])}")
        if omitted > 0:
            parts.append(f"(only the first {len(shown)} of {len(columns)} columns are shown)")
        return "\n".join(parts)

    if len(df):
        sections.append(
            _section(
                "sample",
                _dedupe(
                    [
                        {"label": f"sample rows{label_suffix}", "text": sample_text(_SAMPLE_ROWS_HEAD, _SAMPLE_ROWS_TAIL)},
                        {"label": f"sample rows{label_suffix}", "text": sample_text(3, 0)},
                    ]
                ),
            )
        )
    return sections


def _detail_sections(df: Any, label_suffix: str = "") -> dict[str, Any]:
    """The new per-column/sample detail for a table, as keys to merge into the summary. If building it
    fails for any reason the structural keys (n_rows, columns, dtypes, null_counts), which the
    self-check reads, are still returned exactly as the baseline produced them."""
    try:
        return {"sections": _tabular_sections(df, label_suffix)}
    except Exception as exc:  # noqa: BLE001
        return {"sections": [], "detail_error": f"{type(exc).__name__}: {exc}"}


def _summarize_csv(path: Path) -> dict[str, Any]:
    import pandas as pd

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # mixed-type columns: dtypes are reported as pandas infers them
        df = pd.read_csv(path)
    result: dict[str, Any] = {
        "format": "csv",
        "n_rows": len(df),
        "columns": list(df.columns),
        "dtypes": {c: str(t) for c, t in df.dtypes.items()},
        "null_counts": {c: int(n) for c, n in df.isna().sum().items() if n > 0},
    }
    result.update(_detail_sections(df))
    return result


def _summarize_parquet(path: Path) -> dict[str, Any]:
    import pyarrow.parquet as pq

    pf = pq.ParquetFile(path)
    schema = pf.schema_arrow
    result: dict[str, Any] = {
        "format": "parquet",
        "n_rows": pf.metadata.num_rows,
        "columns": schema.names,
        "dtypes": {name: str(schema.field(name).type) for name in schema.names},
        "sections": [],
    }
    # Everything above is metadata only, as in the baseline. Reading the values is the new part and
    # must never cost the structural keys.
    try:
        suffix = ""
        if path.stat().st_size <= _MAX_TABLE_BYTES:
            df = pf.read().to_pandas()
        else:  # very large: describe the first row group only, and say so
            df = pf.read_row_group(0).to_pandas()
            suffix = " (first row group only)"
        result["null_counts"] = {str(c): int(n) for c, n in df.isna().sum().items() if n > 0}
        result.update(_detail_sections(df, suffix))
    except Exception as exc:  # noqa: BLE001
        result["detail_error"] = f"{type(exc).__name__}: {exc}"
    return result


def _read_tsv(path: Path) -> Any:
    import csv

    import pandas as pd

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            return pd.read_csv(path, sep="\t", low_memory=False)
        except Exception:  # noqa: BLE001 - stray quote characters are common in SEC-style dumps
            return pd.read_csv(path, sep="\t", low_memory=False, quoting=csv.QUOTE_NONE)


def _summarize_tsv(path: Path) -> dict[str, Any]:
    df = _read_tsv(path)
    # `n_rows` is deliberately NOT set for tsv: the self-check's row-parity gate counts inputs that
    # carry it, and its set of tabular inputs is unchanged from the frozen baseline (csv/parquet).
    result: dict[str, Any] = {
        "format": "tsv",
        "n_data_rows": len(df),
        "columns": [str(c) for c in df.columns],
        "dtypes": {str(c): str(t) for c, t in df.dtypes.items()},
        "null_counts": {str(c): int(n) for c, n in df.isna().sum().items() if n > 0},
    }
    result.update(_detail_sections(df))
    return result


# ---------------------------------------------------------------------------------------------
# JSON / JSONL
# ---------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class _Limits:
    depth: int
    list_items: int
    dict_keys: int
    str_chars: int


# Progressively tighter views of the same value; the packer in agent.digest picks the largest that fits.
_JSON_LEVELS = (
    _Limits(10, 40, 100, 500),
    _Limits(6, 12, 40, 200),
    _Limits(4, 5, 25, 100),
    _Limits(2, 3, 20, 60),
)
_NUMERIC_LIST_MIN = 40  # a list of at least this many numbers is described, not listed


def _is_number(x: Any) -> bool:
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def _bound(value: Any, limits: _Limits, depth: int = 0) -> Any:
    """A bounded copy of `value`: deep nesting, long lists, many keys and long strings are replaced
    by explicit '<...>' markers that say how much was left out."""
    if isinstance(value, dict):
        if depth >= limits.depth:
            keys = ", ".join(str(k) for k in list(value)[:6])
            return f"<object with {len(value)} keys: {keys}{', ...' if len(value) > 6 else ''}>"
        out: dict[str, Any] = {}
        for k, v in list(value.items())[: limits.dict_keys]:
            out[str(k)] = _bound(v, limits, depth + 1)
        if len(value) > limits.dict_keys:
            out["<...>"] = f"{len(value) - limits.dict_keys} more keys omitted"
        return out
    if isinstance(value, list):
        n = len(value)
        if n >= _NUMERIC_LIST_MIN and all(_is_number(x) for x in value):
            head = ", ".join(_num(x) for x in value[:6])
            return f"<{n} numbers; first 6: [{head}]; min={_num(min(value))}, max={_num(max(value))}>"
        if depth >= limits.depth:
            return f"<list of {n} items>"
        out_list = [_bound(v, limits, depth + 1) for v in value[: limits.list_items]]
        if n > limits.list_items:
            out_list.append(f"<... {n - limits.list_items} more items omitted>")
        return out_list
    if isinstance(value, str):
        return value if len(value) <= limits.str_chars else value[: limits.str_chars] + "..."
    return value


def _dump(value: Any) -> str:
    return clean_text(json.dumps(value, ensure_ascii=False, default=str))


def _json_variants(value: Any, full_cap: int) -> list[dict[str, str]]:
    caps = (full_cap, 3000, 1200, 500)
    complete = _dump(value)
    if len(complete) <= full_cap:
        variants = [{"label": f"complete value ({len(complete):,} chars)", "text": complete}]
        # A small file is shown whole; also offer a shrunken view for a tight global budget.
        for limits, cap in zip(_JSON_LEVELS[1:], caps[1:]):
            text = _dump(_bound(value, limits))
            if len(text) <= cap and len(text) < len(variants[-1]["text"]):
                variants.append({"label": "bounded view (long lists/strings abbreviated)", "text": text})
        return _dedupe(variants)

    variants = []
    for limits, cap in zip(_JSON_LEVELS, caps):
        text = _dump(_bound(value, limits))
        if len(text) <= cap:
            variants.append(
                {"label": f"bounded view of a {len(complete):,}-char value (long lists/strings abbreviated)", "text": text}
            )
    if not variants:
        text = _dump(_bound(value, _JSON_LEVELS[-1]))
        variants.append(
            {"label": f"bounded view of a {len(complete):,}-char value, cut", "text": text[: caps[-1]] + "...[cut]"}
        )
    return _dedupe(variants)


def _summarize_json(path: Path, referenced: bool) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    cap = REFERENCED_JSON_CAP if referenced else JSON_CAP
    sections = [_section("json_content", _json_variants(data, cap))]
    if isinstance(data, list):
        first_keys = list(data[0].keys()) if data and isinstance(data[0], dict) else None
        return {"format": "json", "shape": "list", "n_items": len(data), "item_keys": first_keys, "sections": sections}
    if isinstance(data, dict):
        return {"format": "json", "shape": "object", "top_level_keys": list(data.keys()), "sections": sections}
    return {"format": "json", "shape": type(data).__name__, "sections": sections}


def _summarize_jsonl(path: Path) -> dict[str, Any]:
    records: list[Any] = []
    n_lines = 0
    n_bad = 0
    key_order: dict[str, None] = {}
    with path.open("r", encoding="utf-8", errors="replace") as handle:
        for line in handle:
            if not line.strip():
                continue
            n_lines += 1
            if len(records) < 50:
                try:
                    record = json.loads(line)
                except ValueError:
                    n_bad += 1
                    continue
                records.append(record)
                if isinstance(record, dict):
                    key_order.update(dict.fromkeys(record))
    result: dict[str, Any] = {"format": "jsonl", "n_records": n_lines, "record_keys": list(key_order)}
    if n_bad:
        result["unparsable_lines_in_first_50"] = n_bad
    if records:
        text = "\n".join(_dump(_bound(r, _JSON_LEVELS[1])) for r in records[:2])
        short = _dump(_bound(records[0], _JSON_LEVELS[2]))
        result["sections"] = [
            _section(
                "sample",
                _dedupe(
                    [
                        {"label": "first 2 records (long strings/lists abbreviated)", "text": text},
                        {"label": "first record (abbreviated)", "text": short},
                    ]
                ),
            )
        ]
    return result


# ---------------------------------------------------------------------------------------------
# Text, HTML, XML
# ---------------------------------------------------------------------------------------------


def _read_text_best_effort(path: Path) -> tuple[str, bool]:
    """(text, was_cut). Binary content (NUL bytes) is reported as (\"\", False) by the caller."""
    size = path.stat().st_size
    with path.open("rb") as handle:
        if size <= _MAX_FULL_READ_BYTES:
            raw = handle.read()
            cut = False
        else:
            head = handle.read(_MAX_FULL_READ_BYTES // 2)
            handle.seek(size - _MAX_FULL_READ_BYTES // 4)
            raw = head + b"\n...[file middle not read]...\n" + handle.read()
            cut = True
    return raw.decode("utf-8", errors="replace"), cut


def _looks_binary(path: Path) -> bool:
    with path.open("rb") as handle:
        return b"\x00" in handle.read(4096)


def _text_variants(text: str, full_cap: int) -> list[dict[str, str]]:
    text = clean_text(text)
    n = len(text)

    def excerpt(cap: int) -> dict[str, str]:
        if n <= cap:
            return {"label": f"complete file ({n:,} chars)", "text": text}
        head = int(cap * 0.75)
        tail = cap - head
        return {
            "label": f"excerpt: first {head:,} and last {tail:,} of {n:,} chars (middle omitted)",
            "text": text[:head] + "\n...[middle omitted]...\n" + text[-tail:],
        }

    return _dedupe([excerpt(c) for c in (full_cap, *_SHRUNK_CAPS) if c <= full_cap])


def _text_sections(path: Path, referenced: bool) -> list[dict[str, Any]] | None:
    if _looks_binary(path):
        return None
    text, _ = _read_text_best_effort(path)
    cap = REFERENCED_TEXT_CAP if referenced else TEXT_CAP
    return [_section("text_content", _text_variants(text, cap))]


def _summarize_text(path: Path, referenced: bool) -> dict[str, Any]:
    suffix = path.suffix.lower()
    result: dict[str, Any] = {"format": suffix.lstrip(".") or "unknown"}
    sections = _text_sections(path, referenced)
    if sections is None:
        result["binary"] = True
        result["sections"] = []
    else:
        result["sections"] = sections
    return result


def _html_to_text(html: str) -> tuple[str, str]:
    """(title, visible text) using only the standard library."""
    from html.parser import HTMLParser

    class _Collector(HTMLParser):
        def __init__(self) -> None:
            super().__init__(convert_charrefs=True)
            self.parts: list[str] = []
            self.title = ""
            self._skip = 0
            self._in_title = False

        def handle_starttag(self, tag: str, attrs: list) -> None:
            if tag in {"script", "style"}:
                self._skip += 1
            elif tag == "title":
                self._in_title = True
            elif tag in {"p", "div", "br", "tr", "li", "h1", "h2", "h3", "h4"}:
                self.parts.append("\n")

        def handle_endtag(self, tag: str) -> None:
            if tag in {"script", "style"} and self._skip:
                self._skip -= 1
            elif tag == "title":
                self._in_title = False

        def handle_data(self, data: str) -> None:
            if self._skip:
                return
            if self._in_title:
                self.title += data
            else:
                self.parts.append(data)

    collector = _Collector()
    collector.feed(html)
    text = re.sub(r"[ \t\r\f\v]+", " ", "".join(collector.parts))
    text = re.sub(r"\n\s*\n+", "\n", text).strip()
    return collector.title.strip(), text


def _summarize_html(path: Path, referenced: bool) -> dict[str, Any]:
    raw, _ = _read_text_best_effort(path)
    title, text = _html_to_text(raw)
    cap = REFERENCED_TEXT_CAP if referenced else 3000
    variants = _text_variants(text, cap)
    for variant in variants:
        variant["label"] = "visible text of the HTML (tags stripped); " + variant["label"]
    result: dict[str, Any] = {"format": "html", "visible_text_chars": len(text), "sections": [_section("text_content", variants)]}
    if title:
        result["title"] = clean_text(title)[:200]
    return result


def _summarize_xml(path: Path, referenced: bool) -> dict[str, Any]:
    import xml.etree.ElementTree as ET
    from collections import Counter

    raw, _ = _read_text_best_effort(path)
    result: dict[str, Any] = {"format": path.suffix.lower().lstrip(".")}
    structure = ""
    if path.stat().st_size <= 5_000_000:
        try:
            root = ET.fromstring(raw)
            counts = Counter(child.tag.split("}")[-1] for child in root.iter())
            top = ", ".join(f"{tag}×{n}" for tag, n in counts.most_common(15))
            structure = f"root element: {root.tag.split('}')[-1]}; element counts: {top}"
            result["root_element"] = root.tag.split("}")[-1]
        except ET.ParseError:
            structure = "(not well-formed XML)"
    cap = REFERENCED_TEXT_CAP if referenced else 4000
    variants = _text_variants(raw, cap)
    sections = []
    if structure:
        sections.append(_section("structure", [{"label": "XML structure", "text": clean_text(structure)}]))
    sections.append(_section("text_content", variants))
    result["sections"] = sections
    return result


# ---------------------------------------------------------------------------------------------
# Workbooks and archives
# ---------------------------------------------------------------------------------------------

_MAX_SHEETS = 6
_MAX_SHEET_ROWS = 8
_MAX_SHEET_COLS = 20
_MAX_ZIP_MEMBERS = 40
_ZIP_TEXT_EXTENSIONS = {".xml", ".txt", ".csv", ".json", ".htm", ".html", ".xsd", ".md", ".tsv"}


def _summarize_xlsx(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"format": "xlsx", "note": "openpyxl is not installed; workbook contents not read", "sections": []}

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
        detail: list[str] = []
        outline: list[str] = []
        for name in sheet_names[:_MAX_SHEETS]:
            sheet = workbook[name]
            dims = f"{sheet.max_row if sheet.max_row is not None else '?'} rows x {sheet.max_column if sheet.max_column is not None else '?'} columns"
            outline.append(f"sheet {name!r}: {dims}")
            lines = []
            for row in sheet.iter_rows(min_row=1, max_row=_MAX_SHEET_ROWS, max_col=_MAX_SHEET_COLS, values_only=True):
                lines.append(" | ".join("" if v is None else _cell(v) for v in row))
            detail.append(f"sheet {name!r} ({dims}); first {len(lines)} rows:\n" + "\n".join(lines))
        if len(sheet_names) > _MAX_SHEETS:
            outline.append(f"(+{len(sheet_names) - _MAX_SHEETS} more sheets)")
    finally:
        workbook.close()
    return {
        "format": "xlsx",
        "sheet_names": sheet_names,
        "sections": [
            _section(
                "structure",
                _dedupe(
                    [
                        {"label": "workbook contents (cell values, first rows of each sheet)", "text": clean_text("\n\n".join(detail))},
                        {"label": "workbook outline", "text": clean_text("\n".join(outline))},
                    ]
                ),
            )
        ],
    }


def _summarize_zip(path: Path) -> dict[str, Any]:
    from collections import Counter

    with zipfile.ZipFile(path) as archive:
        infos = [i for i in archive.infolist() if not i.is_dir()]
        names = [i.filename for i in infos]
        by_ext = Counter(Path(n).suffix.lower() or "(none)" for n in names)
        total = sum(i.file_size for i in infos)
        listing = [f"{i.filename} ({i.file_size:,} bytes)" for i in infos[:_MAX_ZIP_MEMBERS]]
        if len(infos) > _MAX_ZIP_MEMBERS:
            listing.append(f"(+{len(infos) - _MAX_ZIP_MEMBERS} more members)")
        excerpts: list[str] = []
        for info in infos:
            if len(excerpts) >= 3:
                break
            if Path(info.filename).suffix.lower() in _ZIP_TEXT_EXTENSIONS and info.file_size <= 2_000_000:
                with archive.open(info) as member:  # never extracts to disk; reads a few KB only
                    head = member.read(1500).decode("utf-8", errors="replace")
                excerpts.append(f"--- {info.filename} (first {min(1500, info.file_size):,} bytes) ---\n{head}")
    summary = (
        f"{len(infos):,} members, {total:,} bytes uncompressed; by extension: "
        + ", ".join(f"{ext}×{n}" for ext, n in by_ext.most_common(10))
        + "\n"
        + "\n".join(listing)
    )
    outline = f"{len(infos):,} members; by extension: " + ", ".join(f"{ext}×{n}" for ext, n in by_ext.most_common(10))
    variants = [
        {"label": "archive listing and text-member excerpts", "text": clean_text(summary + ("\n\n" + "\n\n".join(excerpts) if excerpts else ""))},
        {"label": "archive listing", "text": clean_text(summary)},
        {"label": "archive outline", "text": clean_text(outline)},
    ]
    return {
        "format": "zip",
        "n_members": len(infos),
        "sections": [_section("structure", _dedupe(variants))],
    }


# ---------------------------------------------------------------------------------------------
# Cross-file relationships between tabular inputs
# ---------------------------------------------------------------------------------------------

_MAX_PAIR_LINES = 12
_MAX_SHARED_COLUMNS_PER_PAIR = 3


def _read_columns(path: Path, columns: list[str]) -> Any:
    import pandas as pd

    suffix = path.suffix.lower()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        if suffix == ".csv":
            return pd.read_csv(path, usecols=columns)
        if suffix in _PARQUET_EXTENSIONS:
            import pyarrow.parquet as pq

            return pq.read_table(path, columns=columns).to_pandas()
        if suffix == ".tsv":
            return _read_tsv(path)[columns]
    raise ValueError(f"not a tabular file: {path.name}")


def cross_file_notes(resolved_inputs: list[ResolvedInput], summaries: list[dict[str, Any]]) -> list[str]:
    """How the tabular inputs relate: shared column names, and for those columns how many values the
    two files have in common. Bounded; returns [] when there is nothing to say; never raises."""
    tabular = [
        (r, s)
        for r, s in zip(resolved_inputs, summaries)
        if s.get("format") in {"csv", "parquet", "tsv"} and s.get("columns") and "summary_error" not in s
    ]
    lines: list[str] = []
    for i in range(len(tabular)):
        for j in range(i + 1, len(tabular)):
            if len(lines) >= _MAX_PAIR_LINES:
                return lines
            (ra, sa), (rb, sb) = tabular[i], tabular[j]
            shared = [c for c in sa["columns"] if c in set(sb["columns"])]
            if not shared:
                continue
            try:
                lines.extend(_pair_lines(ra, sa, rb, sb, shared)[: _MAX_PAIR_LINES - len(lines)])
            except Exception:  # noqa: BLE001 - relationship notes are a courtesy
                continue
    return lines


def _pair_lines(ra: ResolvedInput, sa: dict, rb: ResolvedInput, sb: dict, shared: list[str]) -> list[str]:
    names = f"{ra.resolved_path.name} <-> {rb.resolved_path.name}"
    # Prefer date-like and identifier-like shared columns: they are the usual join keys.
    shared = sorted(shared, key=lambda c: (not re.search(r"date|time|id|ticker|symbol|code|key|name", str(c), re.I), str(c)))
    chosen = shared[:_MAX_SHARED_COLUMNS_PER_PAIR]
    fa, fb = _read_columns(ra.resolved_path, chosen), _read_columns(rb.resolved_path, chosen)
    rows_a = sa.get("n_rows", sa.get("n_data_rows"))
    rows_b = sb.get("n_rows", sb.get("n_data_rows"))
    out = [f"{names}: share {len(shared)} column(s) [{', '.join(_cell(c) for c in shared[:8])}]; rows {rows_a:,} vs {rows_b:,}"]
    for column in chosen:
        va, vb = fa[column].dropna(), fb[column].dropna()
        ua, ub = set(va.astype(str)), set(vb.astype(str))
        same_order = len(va) == len(vb) and bool((va.astype(str).values == vb.astype(str).values).all())
        out.append(
            f"  column {_cell(column)!r}: {len(ua):,} vs {len(ub):,} distinct values, {len(ua & ub):,} in common"
            + ("; identical values in identical row order" if same_order else "")
        )
    return [clean_text(line) for line in out]
